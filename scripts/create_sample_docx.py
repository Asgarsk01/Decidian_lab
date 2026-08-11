from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "sample_srs.docx"
IMAGE = ROOT / "sample_workflow.png"


def main() -> None:
    canvas = Image.new("RGB", (720, 220), "white")
    draw = ImageDraw.Draw(canvas)
    boxes = [(30, 70, 190, 150), (280, 70, 440, 150), (530, 70, 690, 150)]
    labels = ["Staff creates order", "Manager approves", "System captures payment"]
    for box, label in zip(boxes, labels):
        draw.rounded_rectangle(box, radius=10, outline="black", width=3)
        draw.text((box[0] + 10, box[1] + 30), label, fill="black")
    draw.line((190, 110, 280, 110), fill="black", width=3)
    draw.line((440, 110, 530, 110), fill="black", width=3)
    canvas.save(IMAGE)

    document = Document()
    document.add_heading("Order Management", level=1)
    document.add_paragraph("The system must use optimistic locking when updating orders.")
    document.add_heading("Approval Flow", level=2)
    document.add_paragraph("Payment capture must not happen before manager approval.")
    document.add_paragraph("Staff may create an order.", style="List Bullet")
    document.add_paragraph("Manager must approve the order.", style="List Number")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "Allowed action"
    row = table.add_row().cells
    row[0].text = "Manager"
    row[1].text = "Approve order"
    document.add_picture(str(IMAGE))
    document.add_paragraph("Figure 1: Required order approval sequence", style="Caption")
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()

