from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import orjson
from docx import Document as WordDocument
from PIL import Image
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from src.config import Settings
from src.db import Base
from src.models import Decision, Document, DocumentBlock, ImageAnalysis
from src.schemas.decision_candidate_schema import (
    CandidateValidationResult,
    DecisionCandidateBatch,
    DecisionCandidatePayload,
)
from src.services.decision_extractor import extract_decision_candidates, validate_candidates
from src.services.docx_parser import _normalize_blocks
from src.services.image_extractor import extract_all_images
from src.services.packet_builder import build_packets
from src.services.review_service import approve_candidate
from src.services.storage_service import StorageService


class FakeGeminiClient:
    def generate_structured(self, *, model, contents, schema):
        if schema is DecisionCandidateBatch:
            return DecisionCandidateBatch(candidates=[DecisionCandidatePayload(
                decision_type="business_rule",
                decision_statement="The system must not capture payment before manager approval.",
                structured_constraints={
                    "actor": "system", "action": "capture payment",
                    "condition": "manager approval completed",
                    "required_behavior": "wait for approval",
                },
                system_area_hint="Order Management",
                criticality_guess="high",
                confidence_score=0.95,
                confidence_reason="The source uses explicit must-not wording.",
                source_block_ids=["b_004"],
                source_image_ids=[],
                source_excerpt="Payment capture must not happen before manager approval.",
                from_source_type="text",
                requires_human_review=True,
            )]), {"text": "fake candidate response"}
        if schema is CandidateValidationResult:
            return CandidateValidationResult(
                validation_status="passed",
                validation_notes="Directly supported by the cited paragraph.",
                is_supported=True,
            ), {"text": "fake validation response"}
        raise AssertionError(f"Unexpected schema {schema}")


class LocalPipelineTest(unittest.TestCase):
    def test_real_docx_to_human_approved_decision(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            image_path = root / "flow.png"
            Image.new("RGB", (80, 40), "white").save(image_path)
            docx_path = root / "fixture.docx"
            word = WordDocument()
            word.add_heading("Orders", level=1)
            word.add_paragraph("Introductory context.")
            word.add_heading("Approval", level=2)
            word.add_paragraph("Payment capture must not happen before manager approval.")
            word.add_paragraph("Staff may create orders.", style="List Bullet")
            table = word.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Role"
            table.cell(0, 1).text = "Action"
            table.cell(1, 0).text = "Manager"
            table.cell(1, 1).text = "Approve"
            word.add_picture(str(image_path))
            word.add_paragraph("Approval flow", style="Caption")
            word.add_picture(str(image_path))
            word.save(docx_path)

            settings = Settings(
                project_root=root,
                storage_dir=root / "storage",
                database_url=f"sqlite:///{(root / 'test.db').as_posix()}",
                gemini_api_key="test",
                gemini_text_model="fake-text",
                gemini_vision_model="fake-vision",
            )
            storage = StorageService(settings)
            engine = create_engine("sqlite://")
            Base.metadata.create_all(engine)

            with Session(engine) as session:
                session.add(Document(
                    id="doc_test", filename="fixture.docx", file_path=str(docx_path),
                    doc_type="srs", status="uploaded",
                ))
                blocks = _normalize_blocks(WordDocument(str(docx_path)))
                self.assertTrue(any(item["type"] == "heading" for item in blocks))
                self.assertTrue(any(item["type"] == "table" for item in blocks))
                self.assertTrue(any(item["type"] == "list" for item in blocks))
                for block in blocks:
                    session.add(DocumentBlock(
                        document_id="doc_test", block_id=block["block_id"], block_type=block["type"],
                        heading_path=orjson.dumps(block.get("heading_path", [])).decode(),
                        text=block.get("text"), markdown=block.get("markdown"),
                        position_index=block["position_index"], source_json=orjson.dumps(block).decode(),
                    ))
                images, warnings = extract_all_images("doc_test", docx_path, storage, session)
                self.assertEqual(len(images), 2)
                self.assertEqual(images[0]["image_hash"], images[1]["image_hash"])
                for image in images:
                    session.add(ImageAnalysis(
                        document_id="doc_test", image_id=image["image_id"], image_hash=image["image_hash"],
                        vision_provider="fake", vision_model="fake-vision", importance="decorative",
                        image_type="unknown", is_decision_relevant=False, comment="Test image checked.",
                        summary=None, extracted_logic_json="[]", unclear_parts_json="[]",
                        confidence_score=0.5, raw_response_json="{}",
                    ))
                session.flush()

                packets = build_packets("doc_test", storage, session)
                self.assertGreaterEqual(len(packets), 1)
                self.assertEqual(sum(len(p["image_analyses"]) for p in packets), 2)
                client = FakeGeminiClient()
                outputs = extract_decision_candidates("doc_test", storage, session, client, settings)
                self.assertEqual(sum(len(item["candidates"]) for item in outputs), 1)
                report = validate_candidates("doc_test", storage, session, client, settings)
                self.assertEqual(report[0]["final_validation_status"], "passed")
                self.assertEqual(session.scalar(select(func.count()).select_from(Decision)), 0)

                candidate_id = next(
                    item["id"] for output in outputs for item in output["candidates"]
                )
                approve_candidate(session, candidate_id, reviewer="test_reviewer")
                session.flush()
                self.assertEqual(session.scalar(select(func.count()).select_from(Decision)), 1)
                session.commit()

    def test_non_docx_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "DOCX files only"):
            StorageService.validate_docx_filename("requirements.pdf")


if __name__ == "__main__":
    unittest.main()
